#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Organizador Inteligente de Pedidos (uniformes)
Desenvolvido por: Douglas Oliveira | getyourwish10@gmail.com

v1.1.0 (revisado):
- Exportar CSV/Excel voltou a funcionar (asksaveasfilename) e não perde zeros ("07")
- Parser separado da interface, com regex pré-compiladas (muito mais rápido)
- Corrigidos nomes "comidos" (ROBERT G -> ROBER, JOAN 10 -> JOA, - PIETRA -> ETRA, GABRIEL -> tamanho G)
- Baby look: "BABY LOOK M" vira BLM (quando o tamanho BL existe na lista)
- Nada de janela aberta fora da thread principal (evita travamentos do Tkinter)
- Auto-update compara versões de verdade, valida o arquivo baixado e guarda backup
- IA local responde em JSON (mais confiável) e só baixa o modelo de visão quando for usado
- Editor de tamanhos/conversões funcionando, tema escuro completo e salvo
- Destaque de números repetidos, resumo por tamanho na ordem de produção e aba "Resumo" no Excel
"""

from __future__ import annotations

import copy
import csv
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import threading
import time
import urllib.error
import urllib.request
import webbrowser
from collections import Counter
from datetime import datetime
from typing import Callable, Dict, List, Optional, Tuple

import tkinter as tk
from tkinter import ttk, scrolledtext, messagebox, filedialog

__version__ = "1.1.0"
URL_VERSAO_REMOTE = "https://github.com/getyourwish10-byte/ORGANIZADOR_LIDER/blob/main/version.json"

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    OPENPYXL_DISPONIVEL = True
except ImportError:
    OPENPYXL_DISPONIVEL = False

try:
    from PIL import Image, ImageGrab, ImageOps
    PIL_DISPONIVEL = True
except ImportError:
    PIL_DISPONIVEL = False

try:
    import pytesseract
    PYTESSERACT_DISPONIVEL = True
    if not shutil.which("tesseract"):
        for _caminho in (r"C:\Program Files\Tesseract-OCR\tesseract.exe",
                         r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
                         os.path.expandvars(r"%LOCALAPPDATA%\Tesseract-OCR\tesseract.exe")):
            if os.path.isfile(_caminho):
                pytesseract.pytesseract.tesseract_cmd = _caminho
                break
except ImportError:
    PYTESSERACT_DISPONIVEL = False

try:
    import ollama
    OLLAMA_DISPONIVEL = True
except ImportError:
    OLLAMA_DISPONIVEL = False


CONFIG_PATH = os.path.join(os.path.expanduser("~"), ".organizador_pedidos_config.json")
HISTORICO_MAXIMO = 30
COLUNAS = ("NOME", "TAMANHO DE CAMISA", "NÚMERO", "TAMANHO DE CALÇÃO")
COLUNAS_TAMANHO = ("TAMANHO DE CAMISA", "TAMANHO DE CALÇÃO")
PREENCHIMENTO = {"TAMANHO DE CAMISA": "SEM TAMANHO", "NÚMERO": "SEM NÚMERO", "TAMANHO DE CALÇÃO": "SEM CALÇÃO"}
VALORES_VAZIOS = {"", "SEM NOME", "SEM TAMANHO", "SEM NÚMERO", "SEM CALÇÃO"}

ORDEM_PADRAO = [
    "G4", "G3", "G2", "GG", "G", "M", "P",
    "BLG3", "BLG2", "BLGG", "BLG", "BLM", "BLP", "BLPP",
    "13_14_PP", "11_12_GG_INF", "8_10_G_INF", "6_7_M_INF",
    "4_5_P_INF", "2_3_PP_INF", "1_BB", "0_RN",
]

CONVERSOES_PADRAO = {
    "EXG": "G2", "EXGG": "G3", "G1": "GG",
    "BLG1": "BLGG", "BABYLOOK G1": "BLGG",
    "PP": "13_14_PP", "13 A 14 ANOS": "13_14_PP", "13 A 14": "13_14_PP", "13-14 ANOS": "13_14_PP",
    "GG INFANTIL": "11_12_GG_INF", "INFANTIL GG": "11_12_GG_INF", "11 A 12 ANOS": "11_12_GG_INF", "11 A 12": "11_12_GG_INF", "11-12 ANOS": "11_12_GG_INF",
    "G INFANTIL": "8_10_G_INF", "INFANTIL G": "8_10_G_INF", "8 A 10 ANOS": "8_10_G_INF", "8 A 10": "8_10_G_INF", "8-10 ANOS": "8_10_G_INF",
    "M INFANTIL": "6_7_M_INF", "INFANTIL M": "6_7_M_INF", "6 A 7 ANOS": "6_7_M_INF", "6 A 7": "6_7_M_INF", "6-7 ANOS": "6_7_M_INF",
    "P INFANTIL": "4_5_P_INF", "INFANTIL P": "4_5_P_INF", "4 A 5 ANOS": "4_5_P_INF", "4 A 5": "4_5_P_INF", "4-5 ANOS": "4_5_P_INF",
    "PP INFANTIL": "2_3_PP_INF", "INFANTIL PP": "2_3_PP_INF", "2 A 3 ANOS": "2_3_PP_INF", "2 A 3": "2_3_PP_INF", "2-3 ANOS": "2_3_PP_INF",
}

MSG_ERRO_CUDA = (
    "⚠️ Incompatibilidade de driver de vídeo (CUDA) detectada!\n\n"
    "O Ollama tentou usar a placa de vídeo NVIDIA, mas o driver está desatualizado.\n\n"
    "Como resolver:\n"
    "1. Atualize o driver da placa NVIDIA (GeForce Experience).\n"
    "OU\n"
    "2. Force o Ollama a usar só o processador:\n"
    "   - Feche o Ollama na bandeja do Windows (ao lado do relógio).\n"
    "   - Crie a variável de ambiente: OLLAMA_LLM_LIBRARY = cpu_avx2\n"
    "   - Abra o Ollama novamente."
)


def versao_tupla(versao: str) -> Tuple[int, ...]:
    """'1.0.13' -> (1, 0, 13), para comparar versões corretamente (1.0.9 < 1.0.13)."""
    return tuple(int(n) for n in re.findall(r"\d+", str(versao)))


def ollama_exe() -> Optional[str]:
    for caminho in (shutil.which("ollama"), os.path.expandvars(r"%LOCALAPPDATA%\Programs\Ollama\ollama.exe")):
        if caminho and os.path.isfile(caminho):
            return caminho
    return None


def eh_erro_cuda(texto: str) -> bool:
    return any(chave in texto for chave in ("CUDA", "PTX", "status code: 500"))


# ==========================================================================
# Parser (independente da interface: pode ser testado sozinho)
# ==========================================================================

LETRA = "A-Z0-9À-Ÿ"
ANTES = rf"(?<![{LETRA}])"
DEPOIS = rf"(?![{LETRA}])"
SEP = r"[\s\-:/;.,]*"
CALCAO = r"(?:CAL[CÇ][AÃ]O|CALS[AÃ]O|CAUCAO|SHORT[ES]?|BERMUDA)"
CAMISA = r"(?:CAMISETA|CAMISA|CAMIZA|BLUSA)"
VALOR_NUMERO = r"(?:\d+|PI|π)"

RE_IGNORAR = re.compile(r"PATROCÍNIO|NUMERAÇÃO|JOGO|MATERIAL|PENHAROL")
RE_SUBSTITUTO = re.compile(r"^\*?\s*SUB\s*\d+")
RE_ENUMERACAO = re.compile(r"^\s*\d+\s*[-.)]\s*")
RE_SEM_NUMERO = re.compile(r"\bS\s*/\s*N\b")
RE_QUANTIDADE = re.compile(r"^\s*\*?\s*(\d+)\s*(?:X|UN|PCT)\b")
RE_KIT_QTD = re.compile(r"^\s*\*?\s*\d+\s")
RE_BABYLOOK = re.compile(r"\bBABY\s*LOOK\b")
RE_NUM_EXPLICITO = re.compile(rf"{ANTES}(?:N[UÚ]MERO|NUM\.?|N[º°O.]|N)[\s\-:;]*({VALOR_NUMERO}){DEPOIS}|#\s*({VALOR_NUMERO}){DEPOIS}")
RE_NUM_BARRA = re.compile(rf"[/\-]\s*({VALOR_NUMERO}){DEPOIS}")
RE_NUM_SOLTO = re.compile(rf"{ANTES}(\d{{1,3}}|PI|π){DEPOIS}")
RE_PALAVRAS_REMOVER = re.compile(
    r"\b(?:ESCREVER|COM|SEM|NOME|AVULS[AO]S?|CAMISETAS?|CAMISAS?|CAMIZAS?|BLUSAS?|BABY\s*LOOK|ADULTO|INFANTIL|"
    r"KITS?|GOLEIRO|DE|E|CONJUNTOS?|CAL[CÇ][AÃ]O|CALS[AÃ]O|CAUCAO|SHORT[ES]?|BERMUDA|TAMANHO|TAM|T|"
    r"N[UÚ]MERO|NUM|N[º°O]|N)\b\.?|#"
)
RE_PONTAS = re.compile(r"^[\-/:,._–—*\s]+|[\-/:,._–—*\s]+$")
RE_CARACTERES_INVALIDOS = re.compile(r"[^A-Z0-9À-Ÿ'\-\s]")
RE_HIFEN_SOLTO = re.compile(r"(^|\s)[\-']+(\s|$)")

RE_FICHA = re.compile(r"(?:tamanho|nome|n[úu]mero)[/\s]*:", re.IGNORECASE)
RE_SEPARADOR_BLOCOS = re.compile(r"_{3,}|-{3,}")
RE_FICHA_INICIO_TAM = re.compile(r"^\s*(?:TAMANHO|TAM)\b", re.IGNORECASE)
RE_FICHA_TEM_DADO = re.compile(r"NOME|N[ÚU]MERO|\d+", re.IGNORECASE)
RE_FICHA_TAM = re.compile(r"\b(?:TAMANHO|TAM)\b\.?")
RE_FICHA_NUM = re.compile(rf"N[ÚU]MERO[^\d:\n]*:?\s*({VALOR_NUMERO}){DEPOIS}")
RE_FICHA_PREFIXO_NOME = re.compile(r"^(nome(\s+na\s+camisa|\s+da\s+camisa)?)[/\s]*:?\s*", re.IGNORECASE)


def _remover(texto: str, m: re.Match) -> str:
    return f"{texto[:m.start()]} {texto[m.end():]}"


class ParserPedidos:
    """Transforma texto desorganizado em pedidos {NOME, TAMANHO DE CAMISA, NÚMERO, TAMANHO DE CALÇÃO}."""

    def __init__(self, ordem: List[str], conversoes: Dict[str, str]):
        self.ordem = list(ordem)
        self.conversoes = dict(conversoes)
        self.posicao = {t: i for i, t in enumerate(self.ordem)}

        termos = sorted(set(self.ordem) | set(self.conversoes), key=lambda t: (-len(t), t))
        alternativas = "|".join(re.escape(t) for t in termos) or r"(?!)"
        termo = rf"{ANTES}(?P<tam>{alternativas}){DEPOIS}"
        self.re_termo = re.compile(termo)
        self.re_calcao = re.compile(rf"{ANTES}(?:TAMANHO\s+(?:D[AEO]\s+)?{CALCAO}|{CALCAO}\s+TAMANHO|{CALCAO}){SEP}{termo}")
        self.re_camisa = re.compile(
            rf"{ANTES}(?:TAMANHO\s+(?:D[AEO]\s+)?{CAMISA}|{CAMISA}\s+TAMANHO|TAMANHO|TAM\.|TAM|T|{CAMISA}|BABY\s*LOOK){SEP}{termo}"
        )

    # ---------------------------------------------------------------- utilidades
    def converter(self, tamanho: str) -> str:
        if not tamanho or "SEM" in tamanho:
            return ""
        return self.conversoes.get(tamanho, tamanho)

    def _babylook(self, tamanho: str) -> str:
        """G -> BLG, G1 -> BLGG, PP -> BLPP... só se o tamanho BL existir na lista."""
        if not tamanho or tamanho.startswith("BL"):
            return self.converter(tamanho)
        for candidato in ("BL" + tamanho, "BL" + self.converter(tamanho)):
            convertido = self.conversoes.get(candidato, candidato)
            if convertido in self.posicao:
                return convertido
        return self.converter(tamanho)

    def peso(self, pedido: dict) -> Tuple[int, int]:
        camisa, calcao = pedido["TAMANHO DE CAMISA"], pedido["TAMANHO DE CALÇÃO"]
        if camisa in self.posicao:
            return (self.posicao[camisa], 0)
        if calcao in self.posicao:
            return (self.posicao[calcao], 1)
        return (999, 99)

    @staticmethod
    def novo_pedido(nome="", camisa="", numero="", calcao="") -> dict:
        return {"NOME": nome.upper(), "TAMANHO DE CAMISA": camisa, "NÚMERO": numero, "TAMANHO DE CALÇÃO": calcao}

    # ---------------------------------------------------------------- entrada
    def processar(self, texto: str, preencher: bool = True) -> Tuple[List[dict], List[str]]:
        texto = re.sub(r"Número Tabela camisa da equipe:?", "", texto, flags=re.IGNORECASE)
        if RE_FICHA.search(texto) or "___" in texto:
            pedidos, nao_reconhecidas = self._processar_fichas(texto)
        else:
            pedidos, nao_reconhecidas = self._processar_linhas(texto)
        if preencher:
            for p in pedidos:
                for coluna, padrao in PREENCHIMENTO.items():
                    if not p[coluna]:
                        p[coluna] = padrao
        pedidos.sort(key=self.peso)
        return pedidos, nao_reconhecidas

    # ---------------------------------------------------------------- modo lista
    def _processar_linhas(self, texto: str) -> Tuple[List[dict], List[str]]:
        pedidos: List[dict] = []
        nao_reconhecidas: List[str] = []
        contexto_camisa = contexto_calcao = ""

        for original in texto.splitlines():
            original = original.strip()
            if not original:
                continue
            linha = RE_ENUMERACAO.sub("", original.upper())
            linha = RE_SEM_NUMERO.sub(" ", linha)
            if RE_IGNORAR.search(linha) or RE_SUBSTITUTO.match(linha):
                continue

            quantidade = 1
            eh_kit = "KIT" in linha
            eh_conjunto = "CONJUNTO" in linha
            eh_babylook = bool(RE_BABYLOOK.search(linha))
            if eh_kit:
                linha = RE_KIT_QTD.sub(" ", linha, count=1)
            else:
                m = RE_QUANTIDADE.search(linha)
                if m:
                    quantidade = int(m.group(1))
                    linha = _remover(linha, m)
            if "AVULSA" in linha or "AVULSO" in linha:
                contexto_calcao = ""

            tam_calcao = tam_camisa = ""
            m = self.re_calcao.search(linha)
            if m:
                tam_calcao, linha = m.group("tam"), _remover(linha, m)
            m = self.re_camisa.search(linha)
            if m:
                tam_camisa, linha = m.group("tam"), _remover(linha, m)

            soltos = []  # na ordem em que aparecem: o 1º é camisa, o 2º é calção
            while True:
                m = self.re_termo.search(linha)
                if not m:
                    break
                soltos.append(m.group("tam"))
                linha = _remover(linha, m)
            if not tam_camisa and soltos:
                tam_camisa = soltos.pop(0)
            if not tam_calcao and soltos:
                tam_calcao = soltos.pop(0)

            contexto_camisa = tam_camisa or contexto_camisa
            contexto_calcao = tam_calcao or contexto_calcao
            camisa_final = tam_camisa or contexto_camisa
            calcao_final = tam_calcao or contexto_calcao
            if eh_conjunto and not calcao_final and camisa_final:
                calcao_final = camisa_final

            numero = ""
            for regex in (RE_NUM_EXPLICITO, RE_NUM_BARRA, RE_NUM_SOLTO):
                m = regex.search(linha)
                if m:
                    numero = next(g for g in m.groups() if g).upper()
                    linha = _remover(linha, m)
                    break

            linha = RE_PALAVRAS_REMOVER.sub(" ", linha)
            linha = RE_PONTAS.sub("", linha)
            linha = RE_CARACTERES_INVALIDOS.sub(" ", linha)
            nome = re.sub(r"\s+", " ", linha).strip()
            nome = RE_HIFEN_SOLTO.sub(r"\1\2", nome).strip()
            if nome.isdigit():
                nome = ""

            original_up = original.upper()
            if eh_kit:
                eh_pedido = bool(numero or nome)
            else:
                eh_pedido = bool(numero or nome or "CAMISA" in original_up or "BABY" in original_up)
                if eh_pedido and "AVULSAS" in original_up and not tam_camisa:
                    eh_pedido = False

            if not eh_pedido:
                nao_reconhecidas.append(original)
                continue

            camisa = self._babylook(camisa_final) if eh_babylook else self.converter(camisa_final)
            calcao = self.converter(calcao_final)
            numero = "" if "SEM" in numero else numero
            pedidos.extend(self.novo_pedido(nome, camisa, numero, calcao) for _ in range(quantidade))

        return pedidos, nao_reconhecidas

    # ---------------------------------------------------------------- modo ficha
    def _dividir_blocos(self, texto: str) -> List[str]:
        if RE_SEPARADOR_BLOCOS.search(texto):
            return [b.strip() for b in RE_SEPARADOR_BLOCOS.split(texto) if b.strip()]
        blocos, atual = [], []
        for linha in (l.strip() for l in texto.splitlines()):
            if not linha:
                continue
            if RE_FICHA_INICIO_TAM.match(linha) and any(RE_FICHA_TEM_DADO.search(l) for l in atual):
                blocos.append("\n".join(atual))
                atual = [linha]
            else:
                atual.append(linha)
        if atual:
            blocos.append("\n".join(atual))
        return blocos

    def _processar_fichas(self, texto: str) -> Tuple[List[dict], List[str]]:
        pedidos: List[dict] = []
        nao_reconhecidas: List[str] = []

        for bloco in self._dividir_blocos(texto):
            tam_camisa = tam_calcao = nome = numero = ""
            eh_conjunto = "CONJUNTO" in bloco.upper()

            for linha in bloco.splitlines():
                original = linha.strip()
                if not original:
                    continue
                up = original.upper()

                m = RE_FICHA_TAM.search(up)
                if m:
                    resto = up[m.end():]
                    mt = self.re_termo.search(resto) or re.search(r"(?P<tam>[A-Z0-9_]+)", resto)
                    if mt:
                        if re.search(CALCAO, up):
                            tam_calcao = mt.group("tam")
                        else:
                            tam_camisa = mt.group("tam")
                    continue

                m = RE_FICHA_NUM.search(up)
                if m:
                    numero = m.group(1)
                    continue

                if "NOME" in up:
                    partes = original.split(":", 1)
                    if len(partes) > 1 and partes[1].strip():
                        nome = partes[1].strip()
                    continue

                if re.fullmatch(r"\d+|PI|π", up):  # número sozinho na linha
                    numero = numero or up
                    continue

                m = self.re_termo.match(up)  # tamanho sozinho na linha (ex.: "G" ou "EXG")
                if m and not tam_camisa:
                    tam_camisa = m.group("tam")
                    original = original[m.end():].strip()
                    up = original.upper()
                    if not original:
                        continue

                if not nome and not any(k in up for k in ("TAMANHO", "NUMERO", "NÚMERO", "CONJUNTO")):
                    nome = original

            nome = RE_FICHA_PREFIXO_NOME.sub("", nome).strip()
            nome = re.sub(r"[\xa0\t]+", " ", nome).strip()
            camisa = self.converter(tam_camisa)
            calcao = self.converter(tam_calcao) or (camisa if eh_conjunto else "")

            if nome or numero or camisa:
                pedidos.append(self.novo_pedido(nome, camisa, numero, calcao))
            else:
                nao_reconhecidas.append(bloco)

        return pedidos, nao_reconhecidas


# ==========================================================================
# Interface
# ==========================================================================

PALETAS = {
    "claro": {
        "bg": "#F4F6F9", "painel": "#FFFFFF", "texto": "#0F172A", "texto2": "#334155",
        "aviso_bg": "#FFF7ED", "aviso_fg": "#9A3412", "aviso_txt": "#7C2D12",
        "dup_nome": "#FEE2E2", "dup_num": "#FEF3C7", "cabecalho": "#1E293B",
    },
    "escuro": {
        "bg": "#111827", "painel": "#1F2937", "texto": "#E5E7EB", "texto2": "#CBD5E1",
        "aviso_bg": "#2A1F14", "aviso_fg": "#FDBA74", "aviso_txt": "#FED7AA",
        "dup_nome": "#7F1D1D", "dup_num": "#713F12", "cabecalho": "#0B1220",
    },
}


class AplicativoPedidosMagico:
    def __init__(self, root: tk.Tk):
        self.root = root
        self.root.geometry("1280x880")
        self.root.minsize(980, 680)

        self.ordem_tamanhos = list(ORDEM_PADRAO)
        self.conversoes = dict(CONVERSOES_PADRAO)
        self.preencher_padrao = tk.BooleanVar(value=True)
        self.tema_escuro = tk.BooleanVar(value=False)
        self.perguntar_ollama = True
        self.modelo_texto = "llama3"
        self.modelo_visao = "llava"
        self.carregar_config()
        self.parser = ParserPedidos(self.ordem_tamanhos, self.conversoes)

        self.pedidos_atuais: List[dict] = []
        self.linhas_nao_reconhecidas: List[str] = []
        self.historico: List[List[dict]] = []
        self.futuro: List[List[dict]] = []
        self.arquivo_atual: Optional[str] = None
        self._agrupar = True
        self._ordenacao_reversa: Dict[str, bool] = {}
        self._linhas_visiveis: List[Optional[dict]] = []
        self._tarefa_ocupada = False
        self._cores_atuais = PALETAS["claro"]

        self.criar_interface()
        self._configurar_atalhos()
        self._atualizar_titulo()
        if self.tema_escuro.get():
            self.alternar_tema()
        self.root.protocol("WM_DELETE_WINDOW", self._fechar)

        self.root.after(1000, self.verificar_ollama_ao_iniciar)
        self.root.after(2000, lambda: self.verificar_atualizacoes(silencioso=True))

    # ---------------------------------------------------------------- config
    def carregar_config(self) -> None:
        try:
            with open(CONFIG_PATH, encoding="utf-8") as f:
                dados = json.load(f)
        except (OSError, ValueError):
            return
        if not isinstance(dados, dict):
            return
        if isinstance(dados.get("ordem_tamanhos"), list) and dados["ordem_tamanhos"]:
            self.ordem_tamanhos = [str(t) for t in dados["ordem_tamanhos"]]
        if isinstance(dados.get("conversoes"), dict):
            self.conversoes = {str(k): str(v) for k, v in dados["conversoes"].items()}
        self.preencher_padrao.set(bool(dados.get("preencher_padrao", True)))
        self.tema_escuro.set(bool(dados.get("tema_escuro", False)))
        self.perguntar_ollama = bool(dados.get("perguntar_ollama", True))
        self.modelo_texto = dados.get("modelo_texto") or self.modelo_texto
        self.modelo_visao = dados.get("modelo_visao") or self.modelo_visao

    def salvar_config(self) -> None:
        try:
            with open(CONFIG_PATH, "w", encoding="utf-8") as f:
                json.dump({
                    "ordem_tamanhos": self.ordem_tamanhos,
                    "conversoes": self.conversoes,
                    "preencher_padrao": self.preencher_padrao.get(),
                    "tema_escuro": self.tema_escuro.get(),
                    "perguntar_ollama": self.perguntar_ollama,
                    "modelo_texto": self.modelo_texto,
                    "modelo_visao": self.modelo_visao,
                }, f, ensure_ascii=False, indent=2)
        except OSError:
            pass

    def _fechar(self) -> None:
        self.salvar_config()
        self.root.destroy()

    # ---------------------------------------------------------------- threads
    def _na_interface(self, funcao: Callable, *args) -> None:
        """Agenda uma função na thread da interface (Tkinter não pode ser usado de outras threads)."""
        try:
            self.root.after(0, funcao, *args)
        except (RuntimeError, tk.TclError):
            pass  # janela já foi fechada

    def set_status(self, msg: str, progresso: Optional[int] = None) -> None:
        if threading.current_thread() is not threading.main_thread():
            self._na_interface(self.set_status, msg, progresso)
            return
        self.status_var.set(f"{datetime.now():%H:%M:%S}  |  {msg}")
        if progresso is not None:
            self.progress_bar["value"] = progresso

    def _em_segundo_plano(self, alvo: Callable, *args) -> bool:
        """Roda uma tarefa de IA/OCR por vez, com os botões travados enquanto isso."""
        if self._tarefa_ocupada:
            messagebox.showinfo("Aguarde", "Já existe uma leitura/organização em andamento.", parent=self.root)
            return False
        self._tarefa_ocupada = True
        self._travar_botoes_ia(True)

        def executar():
            try:
                alvo(*args)
            finally:
                self._na_interface(self._liberar_tarefa)

        threading.Thread(target=executar, daemon=True).start()
        return True

    def _liberar_tarefa(self) -> None:
        self._tarefa_ocupada = False
        self._travar_botoes_ia(False)

    def _travar_botoes_ia(self, travar: bool) -> None:
        for botao in self.botoes_ia:
            botao.config(state="disabled" if travar else "normal")

    # ---------------------------------------------------------------- atualização
    def verificar_atualizacoes(self, silencioso: bool = True) -> None:
        self.set_status("🔍 Verificando atualizações...", 15)

        def tarefa():
            try:
                req = urllib.request.Request(URL_VERSAO_REMOTE, headers={"User-Agent": "OrganizadorPedidos"})
                with urllib.request.urlopen(req, timeout=8) as resp:
                    dados = json.loads(resp.read().decode("utf-8"))
                self._na_interface(self._resultado_atualizacao, dados, silencioso)
            except urllib.error.HTTPError as e:
                self.set_status(f"⚠️ Erro HTTP {e.code} ao verificar atualizações.", 0)
            except Exception:
                self.set_status("⚠️ Sem conexão para verificar atualizações.", 0)
                if not silencioso:
                    self._na_interface(lambda: messagebox.showerror(
                        "Atualizações", "Não foi possível conectar ao GitHub.", parent=self.root))

        threading.Thread(target=tarefa, daemon=True).start()

    def _resultado_atualizacao(self, dados: dict, silencioso: bool) -> None:
        remota = str(dados.get("version", ""))
        if not remota or versao_tupla(remota) <= versao_tupla(__version__):
            self.set_status(f"✅ Versão {__version__} é a mais recente.", 100)
            if not silencioso:
                messagebox.showinfo("Atualizações", "Seu programa já está na versão mais recente!", parent=self.root)
            return

        self.set_status(f"Nova versão {remota} disponível!", 100)
        url = str(dados.get("url", ""))
        if not url.startswith("https://"):
            messagebox.showerror("Atualização", "O endereço de download da atualização é inválido.", parent=self.root)
            return
        novidades = dados.get("changelog", "Melhorias e correções gerais.")
        if messagebox.askyesno("Atualização disponível",
                               f"Nova versão {remota} disponível (você tem a {__version__}).\n\n"
                               f"Novidades:\n{novidades}\n\nAtualizar agora?", parent=self.root):
            threading.Thread(target=self._baixar_atualizacao, args=(url, dados.get("sha256")), daemon=True).start()

    def _baixar_atualizacao(self, url: str, sha256_esperado: Optional[str]) -> None:
        caminho_atual = os.path.abspath(sys.executable if getattr(sys, "frozen", False) else sys.argv[0])
        caminho_temp = caminho_atual + ".novo"
        try:
            def progresso(bloco, tamanho_bloco, total):
                if total > 0:
                    p = min(90, 10 + bloco * tamanho_bloco * 80 // total)
                    self.set_status(f"Baixando atualização... {p}%", p)

            urllib.request.urlretrieve(url, caminho_temp, reporthook=progresso)
            with open(caminho_temp, "rb") as f:
                conteudo = f.read()

            if sha256_esperado and hashlib.sha256(conteudo).hexdigest().lower() != str(sha256_esperado).lower():
                raise ValueError("o arquivo baixado não confere com o sha256 publicado")
            if caminho_atual.lower().endswith(".py"):
                compile(conteudo, caminho_temp, "exec")  # garante que é Python válido, não uma página de erro
            elif not conteudo.startswith(b"MZ"):
                raise ValueError("o arquivo baixado não é um executável do Windows")
        except Exception as e:
            try:
                os.remove(caminho_temp)
            except OSError:
                pass
            erro = str(e)
            self.set_status("Falha na atualização.", 0)
            self._na_interface(lambda: messagebox.showerror("Atualização", f"Falha ao atualizar:\n{erro}", parent=self.root))
            return
        self._na_interface(self._aplicar_atualizacao, caminho_atual, caminho_temp)

    def _aplicar_atualizacao(self, caminho_atual: str, caminho_temp: str) -> None:
        try:
            if caminho_atual.lower().endswith(".py"):
                shutil.copy2(caminho_atual, caminho_atual + ".bak")
                os.replace(caminho_temp, caminho_atual)
                messagebox.showinfo("Atualizado", "Programa atualizado! Ele será reiniciado agora.", parent=self.root)
                subprocess.Popen([sys.executable, caminho_atual, *sys.argv[1:]])
                self._fechar()
                return

# .exe em uso não pode ser substituído: um .bat espera o programa fechar e troca o arquivo
            bat = os.path.join(tempfile.gettempdir(), "atualizar_organizador.bat")
            with open(bat, "w", encoding="utf-8") as f:
                f.write("@echo off\r\nchcp 65001 > nul\r\nset /a n=0\r\n:tentar\r\nset /a n+=1\r\n"
                        "ping -n 2 127.0.0.1 > nul\r\n"
                        f'move /y "{caminho_temp}" "{caminho_atual}" > nul 2>&1\r\n'
                        "if errorlevel 1 if %n% lss 20 goto tentar\r\n"
                        f'start "" "{caminho_atual}"\r\n'
                        'del "%~f0"\r\n')
            messagebox.showinfo("Atualizando", "O programa será fechado para aplicar a atualização.", parent=self.root)
            subprocess.Popen(["cmd", "/c", bat], creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
            self._fechar()
        except Exception as e:
            messagebox.showerror("Atualização", f"Falha ao aplicar a atualização:\n{e}", parent=self.root)

    # ---------------------------------------------------------------- Ollama
    def verificar_ollama_ao_iniciar(self) -> None:
        if ollama_exe():
            self.set_status("Ollama detectado e pronto para uso local.", 100)
            return
        if sys.platform != "win32" or not self.perguntar_ollama:
            self.set_status("Pronto.", 0)
            return
        if messagebox.askyesno(
                "IA local não encontrada",
                "O motor de IA local (Ollama) não está instalado.\n\n"
                "Instalar agora? O download tem cerca de 1 GB + 4,7 GB do modelo de texto.\n"
                "O modo normal funciona sem ele.", parent=self.root):
            self.instalar_ollama()
        else:
            self.perguntar_ollama = False
            self.salvar_config()
            self.set_status("Você pode instalar a IA depois em Configurações → Instalar IA local.", 0)

    def instalar_ollama(self) -> None:
        if ollama_exe():
            messagebox.showinfo("Ollama", "O Ollama já está instalado.", parent=self.root)
            return
        self._em_segundo_plano(self._instalar_ollama_exec)

    def _instalar_ollama_exec(self) -> None:
        try:
            instalador = os.path.join(tempfile.gettempdir(), "OllamaSetup.exe")

            def progresso(bloco, tamanho_bloco, total):
                if total > 0:
                    p = min(45, 5 + bloco * tamanho_bloco * 40 // total)
                    self.set_status(f"Baixando o Ollama... {p}%", p)

            urllib.request.urlretrieve("https://ollama.com/download/OllamaSetup.exe", instalador, reporthook=progresso)
            self.set_status("Instalando o Ollama...", 50)
            subprocess.run([instalador, "/VERYSILENT", "/SUPPRESSMSGBOXES", "/NORESTART"], check=True)

            exe = ollama_exe()
            if not exe:
                raise RuntimeError("o Ollama foi instalado, mas o executável não foi encontrado")
            self.set_status(f"Baixando o modelo '{self.modelo_texto}' (pode demorar)...", 75)
            if subprocess.run([exe, "pull", self.modelo_texto]).returncode != 0:
                subprocess.Popen([exe, "serve"], creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
                time.sleep(4)
                subprocess.run([exe, "pull", self.modelo_texto], check=True)

            self.set_status("IA local pronta para uso.", 100)
            self._na_interface(lambda: messagebox.showinfo(
                "IA local", "Ollama instalado! O modelo de leitura de imagens será baixado no primeiro uso.",
                parent=self.root))
        except Exception as e:
            erro = str(e)
            self.set_status("Falha na instalação do Ollama.", 0)
            self._na_interface(lambda: messagebox.showerror(
                "Instalação", f"Não foi possível concluir a instalação:\n{erro}", parent=self.root))

    def _ollama_chat(self, modelo: str, mensagens: list, **opcoes):
        """Chama o Ollama; se o modelo não existir ainda, baixa uma vez e tenta de novo."""
        try:
            return ollama.chat(model=modelo, messages=mensagens, **opcoes)
        except ollama.ResponseError as e:
            if getattr(e, "status_code", None) != 404:
                raise
        self.set_status(f"Baixando o modelo '{modelo}' (só na primeira vez, pode demorar)...", 30)
        ollama.pull(modelo)
        return ollama.chat(model=modelo, messages=mensagens, **opcoes)

    def _mostrar_erro_ollama(self, titulo: str, erro: str) -> None:
        if eh_erro_cuda(erro):
            messagebox.showerror("Erro de compatibilidade (Ollama / GPU)", MSG_ERRO_CUDA, parent=self.root)
        else:
            messagebox.showerror(titulo, "Verifique se o aplicativo Ollama está aberto (ícone ao lado do relógio)."
                                         f"\n\nDetalhes:\n{erro}", parent=self.root)

    def _checar_ollama(self) -> bool:
        if OLLAMA_DISPONIVEL:
            return True
        messagebox.showerror("IA local", "A biblioteca 'ollama' não está instalada.\n\nNo terminal, rode:\npip install ollama",
                             parent=self.root)
        return False

    # ---------------------------------------------------------------- interface
    def criar_interface(self) -> None:
        estilo = ttk.Style()
        estilo.theme_use("clam")
        self.estilo = estilo

        menubar = tk.Menu(self.root)
        m_arq = tk.Menu(menubar, tearoff=0)
        m_arq.add_command(label="Abrir lista de um .txt...", command=self.abrir_txt, accelerator="Ctrl+O")
        m_arq.add_command(label="Colar texto da área de transferência", command=self.colar_da_area_transferencia)
        m_arq.add_separator()
        m_arq.add_command(label="🖼️ Ler imagem com IA local (Ollama)...", command=self.ler_imagem_com_ollama)
        m_arq.add_command(label="📋 Colar imagem com IA local", command=self.colar_imagem_com_ollama)
        m_arq.add_command(label="💻 Ler imagem offline (Tesseract)...", command=self.ler_imagem_ocr_offline)
        m_arq.add_command(label="📋 Colar imagem offline (Tesseract)", command=self.colar_imagem_ocr_offline)
        m_arq.add_separator()
        m_arq.add_command(label="Exportar Excel (.xlsx)", command=self.salvar_em_xlsx, accelerator="Ctrl+S")
        m_arq.add_command(label="Exportar CSV (.csv)", command=self.salvar_em_csv)
        m_arq.add_command(label="Copiar para o Excel", command=self.copiar_para_excel, accelerator="Ctrl+Shift+C")
        m_arq.add_separator()
        m_arq.add_command(label="🔄 Verificar atualizações", command=lambda: self.verificar_atualizacoes(silencioso=False))
        m_arq.add_separator()
        m_arq.add_command(label="Sair", command=self._fechar)
        menubar.add_cascade(label="Arquivo", menu=m_arq)

        m_edit = tk.Menu(menubar, tearoff=0)
        m_edit.add_command(label="Desfazer (tabela)", command=self.desfazer, accelerator="Ctrl+Z")
        m_edit.add_command(label="Refazer (tabela)", command=self.refazer, accelerator="Ctrl+Y")
        m_edit.add_separator()
        m_edit.add_command(label="Adicionar linha manual", command=self.adicionar_linha_manual)
        m_edit.add_command(label="Excluir linha(s) selecionada(s)", command=self.excluir_linhas_selecionadas, accelerator="Delete")
        m_edit.add_command(label="Filtrar tabela", command=lambda: self.entrada_filtro.focus_set(), accelerator="Ctrl+F")
        menubar.add_cascade(label="Editar", menu=m_edit)

        m_conf = tk.Menu(menubar, tearoff=0)
        m_conf.add_command(label="Gerenciar tamanhos e conversões...", command=self.abrir_editor_regras)
        m_conf.add_checkbutton(label="Preencher campos vazios (SEM TAMANHO/NÚMERO)", variable=self.preencher_padrao, command=self.salvar_config)
        m_conf.add_checkbutton(label="Tema escuro", variable=self.tema_escuro, command=self.alternar_tema)
        m_conf.add_separator()
        m_conf.add_command(label="Instalar IA local (Ollama)...", command=self.instalar_ollama)
        m_conf.add_command(label="Escolher modelos de IA...", command=self.escolher_modelos)
        m_conf.add_separator()
        m_conf.add_command(label="Restaurar padrões de fábrica", command=self.restaurar_padroes)
        menubar.add_cascade(label="Configurações", menu=m_conf)
        self.root.config(menu=menubar)

        c = PALETAS["claro"]
        self.root.configure(bg=c["bg"])
        self.header_frame = tk.Frame(self.root, bg=c["cabecalho"], pady=15)
        self.header_frame.pack(fill="x")
        tk.Label(self.header_frame, text="ORGANIZADOR INTELIGENTE DE UNIFORMES — TURBO", font=("Segoe UI", 16, "bold"),
                 bg=c["cabecalho"], fg="white").pack()
        tk.Label(self.header_frame, text=f"Versão {__version__} | 100% local com Ollama (texto + visão) e atualização automática",
                 font=("Segoe UI", 9), bg=c["cabecalho"], fg="#94A3B8").pack()

        # Rodapé é empacotado antes da área expansível para nunca ser espremido
        rodape = tk.Frame(self.root, bg="#1E293B", pady=10, padx=15)
        rodape.pack(fill="x", side="bottom")
        barra = tk.Frame(rodape, bg="#1E293B")
        barra.pack(fill="x", pady=(0, 6))
        self.status_var = tk.StringVar(value="🟢 Iniciando...")
        tk.Label(barra, textvariable=self.status_var, font=("Segoe UI", 10, "bold"), bg="#1E293B", fg="#F8FAFC",
                 anchor="w").pack(side="left", fill="x", expand=True)
        self.progress_bar = ttk.Progressbar(barra, orient="horizontal", length=220, mode="determinate", maximum=100)
        self.progress_bar.pack(side="right", padx=(10, 0))
        tk.Label(rodape, text="Desenvolvido por Douglas Oliveira  |  Contato: getyourwish10@gmail.com",
                 font=("Segoe UI", 9), bg="#1E293B", fg="#94A3B8").pack()

        self.frame_entrada = tk.LabelFrame(self.root, text=" Cole a lista desorganizada abaixo (uma por linha ou em fichas): ",
                                           font=("Segoe UI", 10, "bold"), bg=c["bg"], fg=c["texto2"], padx=15, pady=10)
        self.frame_entrada.pack(fill="x", padx=20, pady=(15, 5))

        barra_img = tk.Frame(self.frame_entrada, bg=c["bg"])
        barra_img.pack(fill="x", pady=(0, 6))
        estilo_btn = dict(fg="white", font=("Segoe UI", 9, "bold"), bd=0, cursor="hand2", padx=8, pady=4)
        self.botoes_ia = [
            tk.Button(barra_img, text="🖼️ Ler imagem (IA local)", command=self.ler_imagem_com_ollama, bg="#7C3AED", **estilo_btn),
            tk.Button(barra_img, text="📋🖼️ Colar imagem (IA local)", command=self.colar_imagem_com_ollama, bg="#7C3AED", **estilo_btn),
            tk.Button(barra_img, text="💻 Ler offline (Tesseract)", command=self.ler_imagem_ocr_offline, bg="#475569", **estilo_btn),
        ]
        for botao in self.botoes_ia:
            botao.pack(side="left", padx=(0, 4))

        self.text_area = scrolledtext.ScrolledText(self.frame_entrada, height=7, font=("Segoe UI", 10), bd=1,
                                                   relief="solid", undo=True, bg=c["painel"], fg=c["texto"])
        self.text_area.pack(fill="both", expand=True)

        topo = tk.Frame(self.root, bg=c["bg"])
        topo.pack(fill="x", padx=20, pady=5)
        estilo_grande = dict(fg="white", font=("Segoe UI", 10, "bold"), activeforeground="white", bd=0, cursor="hand2", padx=8, pady=8)
        tk.Button(topo, text="✨ SEPARAR E ORGANIZAR (Ctrl+Enter)", command=self.processar_texto, bg="#2563EB",
                  activebackground="#1D4ED8", **estilo_grande).pack(side="left", padx=(0, 5), ipadx=5, fill="x", expand=True)
        btn_ia = tk.Button(topo, text="🏠 ORGANIZAR COM IA LOCAL (Ollama)", command=self.processar_texto_com_ia_local_thread,
                           bg="#059669", activebackground="#047857", **estilo_grande)
        btn_ia.pack(side="left", padx=(0, 5), ipadx=5, fill="x", expand=True)
        self.botoes_ia.append(btn_ia)
        tk.Button(topo, text="🗑️ Limpar tudo", command=self.limpar_tudo, bg="#64748B", activebackground="#475569",
                  **estilo_grande).pack(side="left", ipadx=5)

        busca = tk.Frame(self.root, bg=c["bg"])
        busca.pack(fill="x", padx=20, pady=(0, 5))
        tk.Label(busca, text="🔎 Filtrar tabela:", bg=c["bg"], fg=c["texto2"], font=("Segoe UI", 9)).pack(side="left")
        self.var_filtro = tk.StringVar()
        self.var_filtro.trace_add("write", lambda *_: self._redesenhar_tabela())
        self.entrada_filtro = tk.Entry(busca, textvariable=self.var_filtro, font=("Segoe UI", 9), width=30,
                                       bg=c["painel"], fg=c["texto"], insertbackground=c["texto"])
        self.entrada_filtro.pack(side="left", padx=5)
        self.lbl_resumo = tk.Label(busca, text="Nenhum pedido gerado ainda.", bg=c["bg"], fg=c["texto2"],
                                   font=("Segoe UI", 9, "bold"), justify="right")
        self.lbl_resumo.pack(side="right")

        acoes = tk.Frame(self.root, bg=c["bg"])
        acoes.pack(fill="x", padx=20, pady=5, side="bottom")
        estilo_acao = dict(fg="white", font=("Segoe UI", 10, "bold"), bd=0, cursor="hand2", padx=15, pady=8)
        tk.Button(acoes, text="📋 COPIAR PARA COLAR NO EXCEL", command=self.copiar_para_excel, bg="#16A34A",
                  **estilo_acao).pack(side="left", padx=5, expand=True, fill="x")
        tk.Button(acoes, text="📊 SALVAR COMO EXCEL (.XLSX)", command=self.salvar_em_xlsx, bg="#0EA5E9",
                  **estilo_acao).pack(side="left", padx=5, expand=True, fill="x")
        tk.Button(acoes, text="💾 SALVAR ARQUIVO (.CSV)", command=self.salvar_em_csv, bg="#D97706",
                  **estilo_acao).pack(side="left", padx=5, expand=True, fill="x")

        self.frame_avisos = tk.LabelFrame(self.root, text=" ⚠️ Linhas não reconhecidas (revise manualmente): ",
                                          font=("Segoe UI", 9, "bold"), bg=c["aviso_bg"], fg=c["aviso_fg"], padx=10, pady=5)
        self.txt_avisos = tk.Text(self.frame_avisos, height=3, font=("Segoe UI", 9), bd=0, bg=c["aviso_bg"],
                                  fg=c["aviso_txt"], wrap="word", state="disabled")
        self.txt_avisos.pack(fill="both", expand=True)
        self._acoes_frame = acoes

        self.frame_tabela = tk.Frame(self.root, bg="#E2E8F0", bd=1, relief="solid")
        self.frame_tabela.pack(fill="both", expand=True, padx=20, pady=5)
        scroll_y = ttk.Scrollbar(self.frame_tabela)
        scroll_y.pack(side="right", fill="y")
        scroll_x = ttk.Scrollbar(self.frame_tabela, orient="horizontal")
        scroll_x.pack(side="bottom", fill="x")
        self.tree = ttk.Treeview(self.frame_tabela, columns=COLUNAS, show="headings", selectmode="extended",
                                 yscrollcommand=scroll_y.set, xscrollcommand=scroll_x.set)
        for col in COLUNAS:
            self.tree.heading(col, text=col, command=lambda c_=col: self.ordenar_por_coluna(c_))
            self.tree.column(col, width=250, anchor="center")
        self.tree.pack(fill="both", expand=True)
        scroll_y.config(command=self.tree.yview)
        scroll_x.config(command=self.tree.xview)
        self.tree.bind("<Double-1>", self.editar_celula)
        self.tree.bind("<Delete>", lambda _e: self.excluir_linhas_selecionadas())
        self._aplicar_estilo_tabela(PALETAS["claro"])

    def _aplicar_estilo_tabela(self, p: dict) -> None:
        self.estilo.configure("Treeview.Heading", font=("Segoe UI", 10, "bold"), background="#1E293B", foreground="white")
        self.estilo.map("Treeview.Heading", background=[("active", "#334155")])
        self.estilo.configure("Treeview", font=("Segoe UI", 10), rowheight=25,
                              background=p["painel"], fieldbackground=p["painel"], foreground=p["texto"])
        self.tree.tag_configure("dup_nome", background=p["dup_nome"])
        self.tree.tag_configure("dup_num", background=p["dup_num"])

    def _configurar_atalhos(self) -> None:
        def so_fora_de_campos(acao):
            # Ctrl+Z/Ctrl+Y dentro da caixa de texto continuam desfazendo o TEXTO, não a tabela
            def handler(_e):
                if isinstance(self.root.focus_get(), (tk.Text, tk.Entry)):
                    return None
                acao()
                return "break"
            return handler

        def ctrl_enter(_e):
            self.processar_texto()
            return "break"

        self.text_area.bind("<Control-Return>", ctrl_enter)
        self.root.bind("<Control-Return>", ctrl_enter)
        self.root.bind("<Control-z>", so_fora_de_campos(self.desfazer))
        self.root.bind("<Control-y>", so_fora_de_campos(self.refazer))
        self.root.bind("<Control-f>", lambda _e: self.entrada_filtro.focus_set())
        self.root.bind("<Control-s>", lambda _e: self.salvar_em_xlsx())
        self.root.bind("<Control-o>", lambda _e: self.abrir_txt())
        self.root.bind("<Control-C>", lambda _e: self.copiar_para_excel())  # Ctrl+Shift+C
        self.entrada_filtro.bind("<Escape>", lambda _e: self.var_filtro.set(""))

    def _atualizar_titulo(self) -> None:
        extra = f" — {os.path.basename(self.arquivo_atual)}" if self.arquivo_atual else ""
        self.root.title(f"Organizador Inteligente de Pedidos - by Douglas Oliveira{extra}")

    def alternar_tema(self) -> None:
        novo = PALETAS["escuro" if self.tema_escuro.get() else "claro"]
        antigo = self._cores_atuais
        mapa = {antigo[k].upper(): novo[k] for k in antigo if k != "cabecalho"}
        mapa[antigo["cabecalho"].upper()] = novo["cabecalho"]

        def trocar(widget):
            for opcao in ("bg", "fg", "insertbackground"):
                try:
                    atual = str(widget.cget(opcao)).upper()
                except tk.TclError:
                    continue
                if atual in mapa:
                    widget.configure(**{opcao: mapa[atual]})
            for filho in widget.winfo_children():
                trocar(filho)

        trocar(self.root)
        self._aplicar_estilo_tabela(novo)
        self._cores_atuais = novo
        self.salvar_config()
        self.set_status("Tema atualizado.")

    # ---------------------------------------------------------------- desfazer
    def _salvar_estado_para_undo(self) -> None:
        self.historico.append(copy.deepcopy(self.pedidos_atuais))
        del self.historico[:-HISTORICO_MAXIMO]
        self.futuro.clear()

    def desfazer(self) -> None:
        if not self.historico:
            self.set_status("Nada para desfazer.")
            return
        self.futuro.append(copy.deepcopy(self.pedidos_atuais))
        self.pedidos_atuais = self.historico.pop()
        self._redesenhar_tabela()
        self.set_status("Ação desfeita.")

    def refazer(self) -> None:
        if not self.futuro:
            self.set_status("Nada para refazer.")
            return
        self.historico.append(copy.deepcopy(self.pedidos_atuais))
        self.pedidos_atuais = self.futuro.pop()
        self._redesenhar_tabela()
        self.set_status("Ação refeita.")

    # ---------------------------------------------------------------- processamento
    def _texto_entrada(self) -> Optional[str]:
        texto = self.text_area.get("1.0", tk.END).strip()
        if not texto:
            messagebox.showwarning("Aviso", "A caixa de texto está vazia!", parent=self.root)
            return None
        return texto

    def _aplicar_resultado(self, pedidos: List[dict], nao_reconhecidas: List[str], origem: str) -> None:
        self._salvar_estado_para_undo()
        self.pedidos_atuais = pedidos
        self.linhas_nao_reconhecidas = nao_reconhecidas
        self._agrupar = True
        self._redesenhar_tabela()
        self._atualizar_avisos()
        extra = f" {len(nao_reconhecidas)} linha(s) não reconhecida(s)." if nao_reconhecidas else ""
        self.set_status(f"Sucesso! {len(pedidos)} pedido(s) organizado(s) ({origem}).{extra}", 100)

    def processar_texto(self) -> None:
        texto = self._texto_entrada()
        if texto is None:
            return
        self.set_status("Processando (modo normal)...", 30)
        pedidos, nao_reconhecidas = self.parser.processar(texto, self.preencher_padrao.get())
        self._aplicar_resultado(pedidos, nao_reconhecidas, "modo normal")

    def processar_texto_com_ia_local_thread(self) -> None:
        texto = self._texto_entrada()
        if texto is None or not self._checar_ollama():
            return
        if self._em_segundo_plano(self._processar_texto_com_ia_local_exec, texto):
            self.set_status("A IA local está organizando os dados...", 40)

    def _prompt_ia(self, texto: str) -> str:
        conversoes = ", ".join(f"{k} → {v}" for k, v in self.conversoes.items())
        return (
            "Você organiza listas de pedidos de uniformes para uma confecção.\n"
            "Transforme o texto abaixo em JSON no formato:\n"
            '{"pedidos": [{"nome": "", "camisa": "", "numero": "", "calcao": ""}]}\n\n'
            "Regras:\n"
            "- Um objeto por peça/pessoa. Se a linha disser uma quantidade (ex.: 3X), repita o objeto.\n"
            "- Nomes em LETRAS MAIÚSCULAS, sem palavras como CAMISA, TAMANHO, NÚMERO.\n"
            f"- Tamanhos válidos: {', '.join(self.ordem_tamanhos)}.\n"
            f"- Converta: {conversoes}.\n"
            "- Se a linha tiver CONJUNTO, o calção tem o mesmo tamanho da camisa.\n"
            "- Número pode ser PI ou π; mantenha como está.\n"
            "- Deixe o campo vazio (\"\") quando a informação não existir. Não invente dados.\n\n"
            f"Texto:\n{texto}"
        )

    def _normalizar_pedido_ia(self, nome, camisa, numero, calcao) -> dict:
        limpo = lambda v: str(v or "").strip().upper()
        pedido = ParserPedidos.novo_pedido(
            limpo(nome), self.parser.converter(limpo(camisa)), limpo(numero), self.parser.converter(limpo(calcao)))
        if pedido["NOME"] == "SEM NOME":
            pedido["NOME"] = ""
        if "SEM" in pedido["NÚMERO"]:
            pedido["NÚMERO"] = ""
        if self.preencher_padrao.get():
            for coluna, padrao in PREENCHIMENTO.items():
                pedido[coluna] = pedido[coluna] or padrao
        return pedido

    def _interpretar_resposta_ia(self, resposta: str) -> List[dict]:
        resposta = resposta.replace("```json", "").replace("```csv", "").replace("```", "").strip()
        pedidos: List[dict] = []
        try:
            dados = json.loads(resposta)
            itens = dados.get("pedidos", []) if isinstance(dados, dict) else dados
            for it in itens if isinstance(itens, list) else []:
                if isinstance(it, dict):
                    pedidos.append(self._normalizar_pedido_ia(it.get("nome"), it.get("camisa"), it.get("numero"), it.get("calcao")))
        except ValueError:  # modelo ignorou o JSON: aceita o formato antigo "NOME; CAMISA; NÚMERO; CALÇÃO"
            for linha in resposta.splitlines():
                partes = [p.strip() for p in linha.split(";")]
                if len(partes) >= 4 and partes[0].upper() != "NOME":
                    pedidos.append(self._normalizar_pedido_ia(*partes[:4]))
        return [p for p in pedidos if any(p[c] not in VALORES_VAZIOS for c in COLUNAS)]

    def _processar_texto_com_ia_local_exec(self, texto: str) -> None:
        try:
            resposta = self._ollama_chat(self.modelo_texto, [{"role": "user", "content": self._prompt_ia(texto)}],
                                         format="json", options={"temperature": 0})
            pedidos = self._interpretar_resposta_ia(resposta["message"]["content"])
        except Exception as e:
            erro = str(e)
            self.set_status("Erro na IA local.", 0)
            self._na_interface(self._mostrar_erro_ollama, "Erro na IA local (Ollama)", erro)
            return
        if not pedidos:
            self.set_status("A IA local não retornou dados válidos.", 0)
            self._na_interface(lambda: messagebox.showwarning(
                "IA local", "A IA não retornou dados válidos. Tente o modo normal.", parent=self.root))
            return
        pedidos.sort(key=self.parser.peso)
        self._na_interface(self._aplicar_resultado, pedidos, [], "IA local")

    # ---------------------------------------------------------------- imagens
    def _imagem_da_area_transferencia(self):
        if not PIL_DISPONIVEL:
            messagebox.showerror("Imagem", "Instale o Pillow: pip install pillow", parent=self.root)
            return None
        try:
            conteudo = ImageGrab.grabclipboard()
        except Exception:
            conteudo = None
        if isinstance(conteudo, list):  # arquivos copiados no Explorer
            caminhos = [c for c in conteudo if c.lower().endswith((".png", ".jpg", ".jpeg", ".bmp", ".webp"))]
            conteudo = Image.open(caminhos[0]) if caminhos else None
        if conteudo is None:
            messagebox.showinfo("Aviso", "Não há imagem copiada na área de transferência.\n"
                                         "Dica: use Win+Shift+S para recortar um print.", parent=self.root)
        return conteudo

    def ler_imagem_com_ollama(self) -> None:
        if not self._checar_ollama():
            return
        caminho = filedialog.askopenfilename(title="Selecione a imagem",
                                             filetypes=[("Imagens", "*.png *.jpg *.jpeg *.bmp *.webp")])
        if caminho:
            self._em_segundo_plano(self._ocr_ollama_exec, caminho)

    def colar_imagem_com_ollama(self) -> None:
        if not self._checar_ollama():
            return
        imagem = self._imagem_da_area_transferencia()
        if imagem is None:
            return
        with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
            imagem.convert("RGB").save(tmp, "PNG")
        self._em_segundo_plano(self._ocr_ollama_exec, tmp.name, True)

    def _ocr_ollama_exec(self, caminho: str, apagar_depois: bool = False) -> None:
        prompt = ("Transcreva EXATAMENTE todo o texto visível nesta imagem, linha por linha, na mesma ordem. "
                  "Não adicione comentários, saudações nem explicações.")
        try:
            self.set_status(f"Lendo imagem com IA local ({self.modelo_visao})...", 50)
            resposta = self._ollama_chat(self.modelo_visao, [{"role": "user", "content": prompt, "images": [caminho]}],
                                         options={"temperature": 0})
            texto = resposta["message"]["content"].strip()
        except Exception as e:
            erro = str(e)
            self.set_status("Erro na leitura de imagem por IA.", 0)
            self._na_interface(self._mostrar_erro_ollama, "Erro de IA local (visão)", erro)
            return
        finally:
            if apagar_depois:
                try:
                    os.remove(caminho)
                except OSError:
                    pass
        if not texto or re.match(r"^(desculpe|sinto muito|i'?m sorry|sorry)", texto, re.IGNORECASE):
            self.set_status("A IA não conseguiu ler esta imagem.", 0)
            self._na_interface(lambda: messagebox.showwarning(
                "IA local", "A IA não conseguiu ler esta imagem. Tente o botão '💻 Ler offline (Tesseract)'.",
                parent=self.root))
            return
        self.set_status("Leitura de imagem concluída.", 100)
        self._na_interface(self._mostrar_janela_revisao_ocr, texto)

    def _aviso_instalar_tesseract(self) -> None:
        if messagebox.askyesno("Tesseract-OCR", "A leitura offline precisa do Tesseract-OCR instalado "
                                                "(marque o idioma Portuguese na instalação).\n\n"
                                                "Abrir a página de download?", parent=self.root):
            webbrowser.open("https://github.com/UB-Mannheim/tesseract/wiki")

    def ler_imagem_ocr_offline(self) -> None:
        if not (PYTESSERACT_DISPONIVEL and PIL_DISPONIVEL):
            self._aviso_instalar_tesseract()
            return
        caminho = filedialog.askopenfilename(filetypes=[("Imagens", "*.png *.jpg *.jpeg *.bmp *.webp")])
        if caminho:
            self._em_segundo_plano(self._ocr_tesseract_exec, Image.open(caminho))

    def colar_imagem_ocr_offline(self) -> None:
        if not (PYTESSERACT_DISPONIVEL and PIL_DISPONIVEL):
            self._aviso_instalar_tesseract()
            return
        imagem = self._imagem_da_area_transferencia()
        if imagem is not None:
            self._em_segundo_plano(self._ocr_tesseract_exec, imagem)

    def _ocr_tesseract_exec(self, imagem) -> None:
        try:
            self.set_status("Lendo imagem offline...", 50)
            # Pré-processamento simples melhora bastante prints de WhatsApp
            img = ImageOps.grayscale(ImageOps.exif_transpose(imagem))
            if img.width < 1600:
                img = img.resize((img.width * 2, img.height * 2), Image.LANCZOS)
            img = ImageOps.autocontrast(img)
            try:
                texto = pytesseract.image_to_string(img, lang="por")
            except pytesseract.TesseractError:  # pacote de português não instalado
                texto = pytesseract.image_to_string(img)
        except pytesseract.TesseractNotFoundError:
            self.set_status("Tesseract não encontrado.", 0)
            self._na_interface(self._aviso_instalar_tesseract)
            return
        except Exception as e:
            erro = str(e)
            self.set_status("Erro na leitura offline.", 0)
            self._na_interface(lambda: messagebox.showerror("OCR", f"Erro na leitura:\n{erro}", parent=self.root))
            return
        self.set_status("Leitura offline concluída.", 100)
        self._na_interface(self._mostrar_janela_revisao_ocr, texto)

    def _mostrar_janela_revisao_ocr(self, texto: str) -> None:
        janela = tk.Toplevel(self.root)
        janela.title("Revisar texto lido")
        janela.geometry("640x500")
        janela.transient(self.root)
        janela.grab_set()
        tk.Label(janela, text="Confira e corrija o texto lido antes de usar:", font=("Segoe UI", 10, "bold")).pack(anchor="w", padx=15, pady=(15, 5))

        botoes = tk.Frame(janela)
        botoes.pack(side="bottom", fill="x", padx=15, pady=15)
        txt = scrolledtext.ScrolledText(janela, font=("Segoe UI", 10), undo=True)
        txt.pack(fill="both", expand=True, padx=15, pady=(0, 5))
        txt.insert(tk.END, texto)

        def usar(substituir: bool, organizar: bool = False):
            valor = txt.get("1.0", tk.END).strip()
            if substituir:
                self.text_area.delete("1.0", tk.END)
                self.text_area.insert(tk.END, valor)
            else:
                self.text_area.insert(tk.END, "\n" + valor)
            janela.destroy()
            self.set_status("Texto da imagem aplicado.")
            if organizar:
                self.processar_texto()

        tk.Button(botoes, text="✨ Usar e organizar", command=lambda: usar(True, True), bg="#2563EB", fg="white",
                  font=("Segoe UI", 9, "bold"), padx=10, pady=6).pack(side="right", padx=5)
        tk.Button(botoes, text="✅ Substituir na caixa principal", command=lambda: usar(True), bg="#475569", fg="white",
                  font=("Segoe UI", 9, "bold"), padx=10, pady=6).pack(side="right", padx=5)
        tk.Button(botoes, text="➕ Adicionar ao final", command=lambda: usar(False), bg="#16A34A", fg="white",
                  font=("Segoe UI", 9, "bold"), padx=10, pady=6).pack(side="right", padx=5)

    # ---------------------------------------------------------------- tabela
    def obter_peso_ordenacao(self, item: dict) -> Tuple[int, int]:
        return self.parser.peso(item)

    def ordenar_por_coluna(self, coluna: str) -> None:
        pos = self.parser.posicao

        def chave_numero(p):
            n = str(p["NÚMERO"]).upper()
            return int(n) if n.isdigit() else (0 if n in ("PI", "π") else 9999)

        chaves = {
            "NOME": lambda p: p["NOME"],
            "TAMANHO DE CAMISA": lambda p: pos.get(p["TAMANHO DE CAMISA"], 999),
            "NÚMERO": chave_numero,
            "TAMANHO DE CALÇÃO": lambda p: pos.get(p["TAMANHO DE CALÇÃO"], 999),
        }
        reverso = self._ordenacao_reversa.get(coluna, False)
        self._salvar_estado_para_undo()
        self.pedidos_atuais.sort(key=chaves[coluna], reverse=reverso)
        self._ordenacao_reversa[coluna] = not reverso
        self._agrupar = coluna in COLUNAS_TAMANHO and not reverso
        self._redesenhar_tabela()
        self.set_status(f"Tabela ordenada por {coluna}{' (decrescente)' if reverso else ''}.")

    def _chave_grupo(self, p: dict) -> str:
        return p["TAMANHO DE CAMISA"] if p["TAMANHO DE CAMISA"] in self.parser.posicao else p["TAMANHO DE CALÇÃO"]

    def _redesenhar_tabela(self) -> None:
        self.tree.delete(*self.tree.get_children())
        filtro = self.var_filtro.get().strip().upper()
        agrupar = self._agrupar and not filtro

        nomes = Counter(p["NOME"].strip() for p in self.pedidos_atuais if p["NOME"].strip() not in VALORES_VAZIOS)
        numeros = Counter(p["NÚMERO"].strip() for p in self.pedidos_atuais if p["NÚMERO"].strip() not in VALORES_VAZIOS)

        self._linhas_visiveis = []
        grupo_atual = None
        for idx, p in enumerate(self.pedidos_atuais):
            if filtro and filtro not in " ".join(p.values()).upper():
                continue
            if agrupar:
                grupo = self._chave_grupo(p)
                if grupo_atual is not None and grupo != grupo_atual:
                    self.tree.insert("", tk.END, values=("", "", "", ""))
                    self._linhas_visiveis.append(None)
                grupo_atual = grupo
            if nomes.get(p["NOME"].strip(), 0) > 1:
                tag = "dup_nome"
            elif numeros.get(p["NÚMERO"].strip(), 0) > 1:
                tag = "dup_num"
            else:
                tag = ""
            self.tree.insert("", tk.END, iid=str(idx), values=[p[c] for c in COLUNAS], tags=(tag,))
            self._linhas_visiveis.append(p)
        self._atualizar_resumo(nomes, numeros)

    def _atualizar_resumo(self, nomes: Counter, numeros: Counter) -> None:
        total = len(self.pedidos_atuais)
        if not total:
            self.lbl_resumo.config(text="Nenhum pedido gerado ainda.")
            return
        pos = self.parser.posicao
        camisas = Counter(p["TAMANHO DE CAMISA"] for p in self.pedidos_atuais if p["TAMANHO DE CAMISA"] not in VALORES_VAZIOS)
        resumo = "  ".join(f"{t}: {q}" for t, q in sorted(camisas.items(), key=lambda kv: pos.get(kv[0], 999)))
        alertas = []
        if sum(1 for q in nomes.values() if q > 1):
            alertas.append("🟥 nomes repetidos")
        if sum(1 for q in numeros.values() if q > 1):
            alertas.append("🟨 números repetidos")
        texto = f"Total: {total}  |  {resumo}"
        if alertas:
            texto += "  |  " + ", ".join(alertas)
        self.lbl_resumo.config(text=texto)

    def _atualizar_avisos(self) -> None:
        if self.linhas_nao_reconhecidas:
            self.frame_avisos.pack(fill="x", padx=20, pady=5, side="bottom", before=self._acoes_frame)
            self.txt_avisos.config(state="normal")
            self.txt_avisos.delete("1.0", tk.END)
            self.txt_avisos.insert(tk.END, "\n".join(f"• {l}" for l in self.linhas_nao_reconhecidas))
            self.txt_avisos.config(state="disabled")
        else:
            self.frame_avisos.pack_forget()

    def editar_celula(self, event) -> None:
        if self.tree.identify_region(event.x, event.y) != "cell":
            return
        item_id = self.tree.identify_row(event.y)
        coluna_id = self.tree.identify_column(event.x)
        if not item_id.isdigit():
            return  # linha separadora
        bbox = self.tree.bbox(item_id, coluna_id)
        if not bbox:
            return
        coluna = COLUNAS[int(coluna_id[1:]) - 1]
        idx = int(item_id)

        entry = tk.Entry(self.tree, font=("Segoe UI", 10), justify="center")
        entry.insert(0, self.pedidos_atuais[idx][coluna])
        entry.select_range(0, tk.END)
        entry.focus_set()
        entry.place(x=bbox[0], y=bbox[1], width=bbox[2], height=bbox[3])
        concluido = False

        def salvar(_e=None):
            nonlocal concluido
            if concluido:
                return
            concluido = True
            valor = entry.get().strip().upper()
            entry.destroy()
            if coluna in COLUNAS_TAMANHO:
                valor = self.parser.conversoes.get(valor, valor)
            if valor == self.pedidos_atuais[idx][coluna]:
                return
            self._salvar_estado_para_undo()
            self.pedidos_atuais[idx][coluna] = valor
            self._redesenhar_tabela()
            self.tree.see(item_id)
            self.set_status(f"Célula atualizada: {coluna} = {valor}")

        def cancelar(_e=None):
            nonlocal concluido
            concluido = True
            entry.destroy()

        entry.bind("<Return>", salvar)
        entry.bind("<Tab>", salvar)
        entry.bind("<FocusOut>", salvar)
        entry.bind("<Escape>", cancelar)

    def adicionar_linha_manual(self) -> None:
        self._salvar_estado_para_undo()
        self.pedidos_atuais.append(ParserPedidos.novo_pedido())
        self._agrupar = False
        self.var_filtro.set("")
        self._redesenhar_tabela()
        ultimo = str(len(self.pedidos_atuais) - 1)
        self.tree.see(ultimo)
        self.tree.selection_set(ultimo)
        self.set_status("Linha em branco adicionada. Dê dois cliques para preencher.")

    def excluir_linhas_selecionadas(self) -> None:
        indices = sorted((int(i) for i in self.tree.selection() if i.isdigit()), reverse=True)
        if not indices:
            return
        self._salvar_estado_para_undo()
        for idx in indices:
            del self.pedidos_atuais[idx]
        self._redesenhar_tabela()
        self.set_status(f"{len(indices)} linha(s) excluída(s). Ctrl+Z desfaz.")

    def limpar_tudo(self) -> None:
        if not messagebox.askyesno("Confirmar", "Limpar o texto e a tabela?", parent=self.root):
            return
        self._salvar_estado_para_undo()
        self.text_area.delete("1.0", tk.END)
        self.pedidos_atuais = []
        self.linhas_nao_reconhecidas = []
        self.var_filtro.set("")
        self._redesenhar_tabela()
        self._atualizar_avisos()
        self.set_status("Tudo limpo. Ctrl+Z recupera a tabela.", 0)

    # ---------------------------------------------------------------- arquivos
    def abrir_txt(self) -> None:
        caminho = filedialog.askopenfilename(filetypes=[("Texto", "*.txt"), ("Todos", "*.*")])
        if not caminho:
            return
        for codificacao in ("utf-8-sig", "cp1252"):  # WhatsApp exporta em UTF-8; Bloco de Notas antigo em ANSI
            try:
                with open(caminho, encoding=codificacao) as f:
                    conteudo = f.read()
                break
            except UnicodeDecodeError:
                continue
            except OSError as e:
                messagebox.showerror("Erro", str(e), parent=self.root)
                return
        self.text_area.delete("1.0", tk.END)
        self.text_area.insert(tk.END, conteudo)
        self.arquivo_atual = caminho
        self._atualizar_titulo()
        self.set_status(f"Arquivo carregado: {os.path.basename(caminho)}")

    def colar_da_area_transferencia(self) -> None:
        try:
            self.text_area.insert(tk.INSERT, self.root.clipboard_get())
            self.set_status("Conteúdo colado.")
        except tk.TclError:
            self.set_status("A área de transferência não contém texto.")

    def _linhas_para_exportar(self) -> List[List[str]]:
        """Exatamente o que está na tela (filtro e separadores), mas a partir dos dados originais."""
        return [[p[c] for c in COLUNAS] if p else ["", "", "", ""] for p in self._linhas_visiveis]

    def _nome_sugerido(self, extensao: str) -> str:
        base = os.path.splitext(os.path.basename(self.arquivo_atual))[0] if self.arquivo_atual else "pedidos"
        return f"{base}_{datetime.now():%Y-%m-%d}{extensao}"

    def _sem_dados(self) -> bool:
        if not self._linhas_visiveis:
            messagebox.showinfo("Exportar", "Não há pedidos na tabela para exportar.", parent=self.root)
            return True
        return False

    def copiar_para_excel(self) -> None:
        if self._sem_dados():
            return
        self.root.clipboard_clear()
        self.root.clipboard_append("\n".join("\t".join(l) for l in self._linhas_para_exportar()))
        self.set_status("✅ Copiado! Cole no Excel com Ctrl+V.", 100)

    def salvar_em_csv(self) -> None:
        if self._sem_dados():
            return
        caminho = filedialog.asksaveasfilename(defaultextension=".csv", initialfile=self._nome_sugerido(".csv"),
                                               filetypes=[("CSV", "*.csv")])
        if not caminho:
            return
        try:
            with open(caminho, "w", newline="", encoding="utf-8-sig") as f:
                w = csv.writer(f, delimiter=";")
                w.writerow(COLUNAS)
                w.writerows(self._linhas_para_exportar())
        except OSError as e:
            messagebox.showerror("Erro", f"Não foi possível salvar (o arquivo está aberto no Excel?):\n{e}", parent=self.root)
            return
        self.set_status(f"CSV salvo: {os.path.basename(caminho)}", 100)

    def salvar_em_xlsx(self) -> None:
        if not OPENPYXL_DISPONIVEL:
            messagebox.showerror("Erro", "Instale o openpyxl: pip install openpyxl", parent=self.root)
            return
        if self._sem_dados():
            return
        caminho = filedialog.asksaveasfilename(defaultextension=".xlsx", initialfile=self._nome_sugerido(".xlsx"),
                                               filetypes=[("Excel", "*.xlsx")])
        if not caminho:
            return
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Pedidos"
            self._planilha(ws, COLUNAS, self._linhas_para_exportar(), [34, 22, 12, 22])
            for celula in ws["C"][1:]:
                celula.number_format = "@"  # número como texto: "07" continua "07"

            pos = self.parser.posicao
            camisas = Counter(p["TAMANHO DE CAMISA"] for p in self.pedidos_atuais if p["TAMANHO DE CAMISA"] not in VALORES_VAZIOS)
            calcoes = Counter(p["TAMANHO DE CALÇÃO"] for p in self.pedidos_atuais if p["TAMANHO DE CALÇÃO"] not in VALORES_VAZIOS)
            tamanhos = sorted(set(camisas) | set(calcoes), key=lambda t: (pos.get(t, 999), t))
            linhas = [[t, camisas.get(t, 0), calcoes.get(t, 0)] for t in tamanhos]
            linhas.append(["TOTAL", sum(camisas.values()), sum(calcoes.values())])
            ws_resumo = wb.create_sheet("Resumo")
            self._planilha(ws_resumo, ("TAMANHO", "CAMISAS", "CALÇÕES"), linhas, [22, 12, 12])
            for celula in ws_resumo[ws_resumo.max_row]:
                celula.font = Font(bold=True)
            wb.save(caminho)
        except OSError as e:
            messagebox.showerror("Erro", f"Não foi possível salvar (o arquivo está aberto no Excel?):\n{e}", parent=self.root)
            return
        self.set_status(f"Planilha salva: {os.path.basename(caminho)}", 100)
        if messagebox.askyesno("Excel salvo", "Planilha salva com as abas Pedidos e Resumo.\nAbrir agora?", parent=self.root):
            try:
                os.startfile(caminho)
            except (AttributeError, OSError):
                pass

    @staticmethod
    def _planilha(ws, cabecalho, linhas, larguras) -> None:
        ws.append(list(cabecalho))
        for linha in linhas:
            ws.append(list(linha))
        for celula in ws[1]:
            celula.font = Font(bold=True, color="FFFFFF")
            celula.fill = PatternFill("solid", fgColor="1E293B")
        centro = Alignment(horizontal="center", vertical="center")
        for linha in ws.iter_rows():
            for celula in linha:
                celula.alignment = centro
        for i, largura in enumerate(larguras):
            ws.column_dimensions[chr(ord("A") + i)].width = largura
        ws.freeze_panes = "A2"

    # ---------------------------------------------------------------- regras
    def _reconstruir_parser(self) -> None:
        self.parser = ParserPedidos(self.ordem_tamanhos, self.conversoes)
        self.salvar_config()
        self._redesenhar_tabela()

    def abrir_editor_regras(self) -> None:
        janela = tk.Toplevel(self.root)
        janela.title("Tamanhos e conversões")
        janela.geometry("760x560")
        janela.transient(self.root)
        janela.grab_set()

        corpo = tk.Frame(janela, padx=15, pady=10)
        corpo.pack(fill="both", expand=True)
        corpo.columnconfigure(0, weight=1)
        corpo.columnconfigure(1, weight=2)
        corpo.rowconfigure(1, weight=1)

        tk.Label(corpo, text="Ordem dos tamanhos\n(um por linha, na ordem da produção)", font=("Segoe UI", 9, "bold"),
                 justify="left").grid(row=0, column=0, sticky="w")
        tk.Label(corpo, text="Conversões\n(como o cliente escreve = tamanho da lista)", font=("Segoe UI", 9, "bold"),
                 justify="left").grid(row=0, column=1, sticky="w", padx=(10, 0))
        txt_ordem = scrolledtext.ScrolledText(corpo, width=18, font=("Consolas", 10))
        txt_ordem.grid(row=1, column=0, sticky="nsew")
        txt_ordem.insert("1.0", "\n".join(self.ordem_tamanhos))
        txt_conv = scrolledtext.ScrolledText(corpo, font=("Consolas", 10))
        txt_conv.grid(row=1, column=1, sticky="nsew", padx=(10, 0))
        txt_conv.insert("1.0", "\n".join(f"{k} = {v}" for k, v in self.conversoes.items()))

        def salvar():
            ordem = list(dict.fromkeys(l.strip().upper() for l in txt_ordem.get("1.0", tk.END).splitlines() if l.strip()))
            if not ordem:
                messagebox.showwarning("Regras", "A lista de tamanhos não pode ficar vazia.", parent=janela)
                return
            conversoes = {}
            for n, linha in enumerate(txt_conv.get("1.0", tk.END).splitlines(), 1):
                if not linha.strip():
                    continue
                de, sep, para = linha.partition("=")
                if not sep or not de.strip() or not para.strip():
                    messagebox.showwarning("Regras", f"Linha {n} das conversões está inválida:\n{linha}\n\nUse: ESCRITO = TAMANHO",
                                           parent=janela)
                    return
                conversoes[de.strip().upper()] = para.strip().upper()
            desconhecidos = sorted({v for v in conversoes.values() if v not in ordem})
            if desconhecidos and not messagebox.askyesno(
                    "Regras", "Estes destinos não estão na lista de tamanhos:\n" + ", ".join(desconhecidos) +
                    "\n\nSalvar mesmo assim?", parent=janela):
                return
            self.ordem_tamanhos, self.conversoes = ordem, conversoes
            self._reconstruir_parser()
            janela.destroy()
            self.set_status("Regras de tamanho salvas. Clique em Separar e Organizar para reaplicar.")

        botoes = tk.Frame(janela, padx=15, pady=10)
        botoes.pack(fill="x")
        tk.Button(botoes, text="💾 Salvar", command=salvar, bg="#2563EB", fg="white", font=("Segoe UI", 9, "bold"),
                  padx=14, pady=5).pack(side="right")
        tk.Button(botoes, text="Cancelar", command=janela.destroy, padx=14, pady=5).pack(side="right", padx=8)

    def escolher_modelos(self) -> None:
        janela = tk.Toplevel(self.root)
        janela.title("Modelos de IA")
        janela.transient(self.root)
        janela.grab_set()
        janela.resizable(False, False)
        var_texto = tk.StringVar(value=self.modelo_texto)
        var_visao = tk.StringVar(value=self.modelo_visao)
        for linha, (rotulo, var, opcoes) in enumerate((
                ("Modelo de texto:", var_texto, ["llama3", "llama3.1", "qwen2.5", "gemma2"]),
                ("Modelo de visão:", var_visao, ["llava", "llama3.2-vision", "qwen2.5vl", "minicpm-v"]))):
            tk.Label(janela, text=rotulo, font=("Segoe UI", 9, "bold")).grid(row=linha, column=0, sticky="w", padx=12, pady=8)
            ttk.Combobox(janela, textvariable=var, values=opcoes, width=24).grid(row=linha, column=1, padx=12, pady=8)

        def salvar():
            self.modelo_texto = var_texto.get().strip() or "llama3"
            self.modelo_visao = var_visao.get().strip() or "llava"
            self.salvar_config()
            janela.destroy()
            self.set_status(f"Modelos: texto = {self.modelo_texto}, visão = {self.modelo_visao}.")

        tk.Button(janela, text="Salvar", command=salvar, bg="#2563EB", fg="white", padx=14).grid(row=2, column=1, sticky="e", padx=12, pady=10)

    def restaurar_padroes(self) -> None:
        if not messagebox.askyesno("Restaurar", "Voltar tamanhos e conversões aos padrões de fábrica?", parent=self.root):
            return
        self.ordem_tamanhos = list(ORDEM_PADRAO)
        self.conversoes = dict(CONVERSOES_PADRAO)
        self._reconstruir_parser()
        self.set_status("Padrões de fábrica restaurados.", 0)


if __name__ == "__main__":
    root = tk.Tk()
    app = AplicativoPedidosMagico(root)
    root.mainloop()
